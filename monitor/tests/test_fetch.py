from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook
import requests

from monitor.fetch import extract_document_text, fetch_source, html_to_text
from monitor.schema import Source


FIXTURE = Path(__file__).parent / "fixtures/recruitment.html"


class FakeResponse:
    status_code = 200
    content = FIXTURE.read_bytes()
    text = content.decode("utf-8")
    headers = {"content-type": "text/html; charset=utf-8"}
    url = "https://example.com/careers/"

    def raise_for_status(self):
        return None


class FakeSession:
    def get(self, url, **kwargs):
        assert kwargs["timeout"] == (10, 30)
        return FakeResponse()


class ForbiddenResponse(FakeResponse):
    status_code = 403

    def raise_for_status(self):
        raise requests.HTTPError("403 Client Error", response=self)


class ForbiddenSession:
    def get(self, url, **kwargs):
        return ForbiddenResponse()


class RssResponse(FakeResponse):
    content = b"""<?xml version='1.0' encoding='UTF-8'?>
    <rss><channel><title>official public articles</title>
    <item><title>2027 geology recruitment</title>
    <link>https://mp.weixin.qq.com/s/example</link></item>
    </channel></rss>"""
    text = content.decode("utf-8")
    headers = {"content-type": "application/rss+xml; charset=utf-8"}
    url = "https://example.com/search.xml"


class RssSession:
    def get(self, url, **kwargs):
        return RssResponse()


class ApiResponse:
    status_code = 200
    headers = {"content-type": "application/json; charset=utf-8"}
    url = "https://gp-api.iguopin.com/api/jobs/v1/list"

    def raise_for_status(self):
        return None

    def json(self):
        return {
            "code": 200,
            "data": {
                "total": 2,
                "list": [
                    {
                        "job_id": "2027-job",
                        "job_name": "矿山地质岗",
                        "company_name": "星海矿业有限公司",
                        "recruitment_type_cn": "校园招聘",
                        "education_cn": "硕士",
                        "major_cn": ["地质工程", "资源勘查工程"],
                        "start_time": "2026-07-14 09:00:00",
                        "end_time": "2026-08-01 23:59:59",
                        "district_list": [{"area_cn": "青海-西宁"}],
                        "contents": "面向2027届高校毕业生，招聘地质工程、资源勘查工程专业硕士",
                    },
                    {
                        "job_id": "social-job",
                        "job_name": "高级地质工程师",
                        "company_name": "成熟人才公司",
                        "recruitment_type_cn": "社会招聘",
                        "education_cn": "本科",
                        "major_cn": ["地质工程"],
                        "contents": "要求十年工作经验",
                    },
                ],
            },
        }


class ApiSession:
    def post(self, url, **kwargs):
        assert url == "https://gp-api.iguopin.com/api/jobs/v1/list"
        assert kwargs["json"]["keyword"] in {"地质学", "地质工程"}
        assert kwargs["headers"]["Subsite"] == "cujiuye"
        return ApiResponse()


def make_source(mode="html"):
    return Source(
        source_id="example",
        name="示例官方招聘",
        url="https://example.com/careers/",
        source_type="集团官网公告",
        trust="official",
        unit_ids=("example",),
        mode=mode,
        official_domains=("example.com",),
    )


def test_html_to_text_drops_navigation_and_scripts():
    text = html_to_text(FIXTURE.read_text(encoding="utf-8"))
    assert "地质学" in text
    assert "window.__STATE__" not in text
    assert "网站导航" not in text


def test_fetch_source_returns_absolute_candidate_links():
    result = fetch_source(make_source(), FakeSession())
    assert result.status == "success"
    assert "地质学" in result.text
    assert result.links == (
        "https://example.com/notice/2027-geology.html",
        "https://example.com/files/jobs.xlsx",
    )


def test_restricted_source_does_not_attempt_network():
    result = fetch_source(make_source(mode="restricted"), None)
    assert result.status == "restricted"
    assert result.text == ""


def test_access_control_is_reported_as_restricted():
    result = fetch_source(make_source(), ForbiddenSession())
    assert result.status == "restricted"


def test_docx_and_xlsx_text_is_extractable():
    document_buffer = BytesIO()
    document = Document()
    document.add_paragraph("地质学硕士")
    document.save(document_buffer)

    workbook_buffer = BytesIO()
    workbook = Workbook()
    workbook.active.append(["岗位", "专业"])
    workbook.active.append(["矿山地质", "地质学"])
    workbook.save(workbook_buffer)

    assert "地质学硕士" in extract_document_text(
        document_buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "https://example.com/jobs.docx",
    )
    assert "矿山地质 地质学" in extract_document_text(
        workbook_buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "https://example.com/jobs.xlsx",
    )


def test_rss_xml_exposes_public_article_links():
    rss_source = Source(
        source_id="rss",
        name="公开文章搜索",
        url="https://example.com/search.xml",
        source_type="公开搜索",
        trust="clue",
        unit_ids=(),
        mode="html",
        official_domains=("example.com",),
        tier="l3",
        role="discovery",
        follow_domains=("example.com", "mp.weixin.qq.com"),
    )
    result = fetch_source(rss_source, RssSession())
    assert result.status == "success"
    assert "2027 geology recruitment" in result.text
    assert "https://mp.weixin.qq.com/s/example" in result.links


def test_iguopin_api_returns_separate_2027_documents_only():
    api_source = Source(
        source_id="sasac-iguopin-jobs",
        name="国资央企招聘平台",
        url="https://cujiuye.iguopin.com/",
        source_type="国资央企权威平台",
        trust="authoritative",
        unit_ids=(),
        mode="iguopin-api",
        official_domains=("cujiuye.iguopin.com", "iguopin.com"),
        tier="l2",
        role="discovery",
        follow_domains=("cujiuye.iguopin.com",),
        query_terms=("地质学", "地质工程"),
    )
    result = fetch_source(api_source, ApiSession())
    assert result.status == "success"
    assert len(result.documents) == 1
    detail_url, text = result.documents[0]
    assert detail_url == "https://cujiuye.iguopin.com/job/detail?id=2027-job"
    assert "星海矿业有限公司2027届校园招聘" in text
    assert "专业要求：地质工程、资源勘查工程" in text
    assert "成熟人才公司" not in text
